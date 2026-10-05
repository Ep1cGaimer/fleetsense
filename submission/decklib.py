"""Build the SIH 2026 idea deck (PS SIH26117) on the OFFICIAL SIH template.

    python submission/build_deck.py            # -> submission/SIH26117_Idea.pptx
    python submission/build_deck.py --team-name "X" --team-id "12345"

Keeps every template element (title placeholder, team-name oval, SIH logo,
footer bar, slide numbers) and the template's "idea details pointers" as the
section headings, fills the body with native shapes/text (no text baked into
images, so the exported PDF is fully machine-readable), drops the template's
instruction slide (6 slides total), and sets document metadata.

Export to PDF with PowerPoint (see submission/README.md) — the portal accepts
PDF only.
"""
from __future__ import annotations

import argparse
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
import sys as _sys
_sys.path.insert(0, str(Path(__file__).resolve().parent))

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "assets" / "SIH2026-IDEA-Presentation-Format.pptx"
SCREENSHOT = HERE / "assets" / "prototype_ui.png"
VERIFIED_ON = "30 Sep 2026"
RUN_TIME = "5–62 s"
OUT = HERE / "SIH26117_Idea.pptx"

# palette: SIH template blue/navy + saffron accent, green = "blocked / safe"
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x00, 0x70, 0xC0)
TINT = RGBColor(0xEA, 0xF2, 0xFB)
SAFFRON = RGBColor(0xE4, 0x6C, 0x0A)
SAFF_TINT = RGBColor(0xFD, 0xF0, 0xE4)
GREEN = RGBColor(0x1E, 0x7B, 0x34)
GREEN_TINT = RGBColor(0xE6, 0xF4, 0xEA)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x5F, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC5, 0xD3, 0xE3)
FONT = "Arial"

IDEA_TITLE = "SOVEREIGN AGENT: AIR-GAPPED AGENTIC AI WORKBENCH"
REPO_URL = "https://github.com/Skr-suyash/Sovereign-Agent"
TESTS = 102


# ---------------------------------------------------------------- helpers

def _runs(p, parts, size, color, bold=False):
    """parts: str | list[str | (text, {bold,color,size,italic})]."""
    if isinstance(parts, str):
        parts = [parts]
    for part in parts:
        text, style = (part, {}) if isinstance(part, str) else part
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = style.get("font", FONT)
        f.size = Pt(style.get("size", size))
        f.bold = style.get("bold", bold)
        f.italic = style.get("italic", False)
        f.color.rgb = style.get("color", color)


def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullet=False,
         space_after=3, margin=0.02, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        _runs(p, para, size, color, bold)
        if bullet:
            _bullet(p)
    return tb


def _bullet(p, char="•", indent_emu=171450):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(indent_emu))
    pPr.set("indent", str(-indent_emu))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bf = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    bc = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(bf)
    pPr.append(bc)


def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.08, dash=False, line_w=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # MSO_LINE.DASH
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def label_box(slide, x, y, w, h, paras, fill=TINT, line=None, size=11,
              color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
              radius=0.12, margin=0.06, bullet=False, space_after=2):
    s = box(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        _runs(p, para, size, color, bold)
        if bullet:
            _bullet(p)
    return s


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def heading(slide, x, y, w, pointer, size=15, color=BLUE):
    """Template 'idea details pointer' rendered as a section heading."""
    return text(slide, x, y, w, 0.34, [[("▸ ", {"color": SAFFRON}), pointer]],
                size=size, color=color, bold=True)


def set_title(slide, title, size=None, box=None):
    t = slide.shapes.title
    if box:
        t.left, t.top, t.width, t.height = (Inches(v) for v in box)
        t.text_frame.word_wrap = True
    runs = t.text_frame.paragraphs[0].runs
    runs[0].text = title
    for r in runs[1:]:
        r.text = ""
    if size:
        runs[0].font.size = Pt(size)


def set_team(slide, team):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Your Team Name":
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = team
            for extra in r[1:]:
                extra.text = ""
            r[0].font.size = Pt(10.5 if len(team) > 12 else 12)
            r[0].font.bold = True
            sh.text_frame.word_wrap = False
            for side in ("margin_left", "margin_right"):
                setattr(sh.text_frame, side, 0)


def drop(slide, name_startswith):
    for sh in list(slide.shapes):
        if sh.name.startswith(name_startswith):
            sh._element.getparent().remove(sh._element)


def picture(slide, path, x, y, w=None, h=None, alt=""):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    return pic


def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sld = sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)



