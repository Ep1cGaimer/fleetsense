"""Slide bodies for the SIH26117 idea deck (slides 2-6).

Layout system: dense bordered cards with centred blue headings, blue header
bars, icon tiles, technology chips, a feature hub, a comparison table and big
stat numbers, on the official SIH template. Every word stays real text (no
text baked into images) so the exported PDF is machine-readable; pictures
carry alt text and icons are marked decorative.

Icons are Lucide (ISC licence): SVG sources in assets/icons/svg, rendered once
to PNG tiles in assets/icons (needs Playwright only when a tile is missing).
"""
from __future__ import annotations

import atexit
import importlib
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ICONS = HERE / "assets" / "icons"

F = "Segoe UI"
INK = RGBColor(0x11, 0x11, 0x11)
INK2 = RGBColor(0x3F, 0x46, 0x50)
INK3 = RGBColor(0x6B, 0x72, 0x80)
BLUE = RGBColor(0x00, 0x70, 0xC0)
BLUE_SOFT = RGBColor(0xDB, 0xEA, 0xFE)
RULE = RGBColor(0xBF, 0xD4, 0xEE)
BORDER = RGBColor(0x1F, 0x1F, 0x1F)
CHIP = RGBColor(0x8A, 0x8F, 0x98)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PINK = RGBColor(0xF0, 0x3E, 0x6E)
PINK_SOFT = RGBColor(0xFF, 0xD9, 0xE2)
GREEN = RGBColor(0x2F, 0xB3, 0x44)
GREEN_DK = RGBColor(0x1E, 0x7B, 0x34)
GREEN_SOFT = RGBColor(0xD9, 0xF7, 0xC9)
RED = RGBColor(0xD9, 0x2D, 0x20)
AMBER = RGBColor(0xB5, 0x6A, 0x00)
DARK = RGBColor(0x1F, 0x29, 0x37)

# icon tile palette: name -> (tile fill, stroke)
PAL = {
    "blue": ("#DBEAFE", "#1D4ED8"), "green": ("#DCFCE7", "#15803D"),
    "orange": ("#FFEDD5", "#C2410C"), "pink": ("#FCE7F3", "#BE185D"),
    "purple": ("#EDE9FE", "#6D28D9"), "yellow": ("#FEF9C3", "#A16207"),
    "teal": ("#CCFBF1", "#0F766E"), "red": ("#FEE2E2", "#B91C1C"),
    "slate": ("#E5E7EB", "#374151"),
}


def _b():
    return importlib.import_module("build_deck")


# ---------------------------------------------------------------- icons

_pw = {}


def _page():
    if not _pw:
        from playwright.sync_api import sync_playwright
        _pw["ctx"] = sync_playwright().start()
        _pw["br"] = _pw["ctx"].chromium.launch(channel="msedge")
        _pw["pg"] = _pw["br"].new_page(device_scale_factor=2)
        atexit.register(lambda: (_pw["br"].close(), _pw["ctx"].stop()))
    return _pw["pg"]


def icon_png(name: str, color: str, tile: bool = True) -> Path:
    out = ICONS / f"{name}-{color}{'' if tile else '-line'}.png"
    if out.exists():
        return out
    svg = (ICONS / "svg" / f"{name}.svg").read_text(encoding="utf-8")
    svg = svg[svg.index("<svg"):]
    bg, fg = PAL[color]
    pg = _page()
    pad, rad = (26, 30) if tile else (0, 0)
    pg.set_content(
        f'<body style="margin:0;background:transparent"><div id="t" style="width:{108 + 2 * pad}px;'
        f'height:{108 + 2 * pad}px;border-radius:{rad}px;background:{bg if tile else "transparent"};'
        f'display:flex;align-items:center;justify-content:center;color:{fg}">{svg}</div>'
        '<style>svg{width:108px;height:108px;stroke-width:1.9}</style></body>')
    pg.locator("#t").screenshot(path=str(out), omit_background=True)
    return out


def ico(slide, name, color, x, y, size, tile=True):
    pic = slide.shapes.add_picture(str(icon_png(name, color, tile)), Inches(x), Inches(y),
                                   Inches(size), Inches(size))
    pic._element.nvPicPr.cNvPr.set("descr", "")  # decorative
    return pic


# ---------------------------------------------------------------- text

def md(s, **base):
    """'plain **bold** plain' -> run parts."""
    return [(seg, {**base, "bold": i % 2 == 1 or base.get("bold", False)})
            for i, seg in enumerate(s.split("**")) if seg]


def TX(slide, x, y, w, h, paras, size=10.5, color=INK, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, bullet=None, space_after=3, line=None, margin=0.0):
    """paras: str (with **bold**) | list of str | list of run-part lists.
    Run part: (text, {bold, italic, color, size, link, font})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        for text, st in (md(para) if isinstance(para, str) else para):
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = st.get("font", F)
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", False)
            f.italic = st.get("italic", False)
            f.color.rgb = st.get("color", color)
            if st.get("link"):
                r.hyperlink.address = st["link"]
        if bullet is not None:
            _bullet(p, bullet)
    return tb


def _bullet(p, color, char="●", indent=0.17):
    pPr = p._p.get_or_add_pPr()
    emu = int(indent * 914400)
    pPr.set("marL", str(emu))
    pPr.set("indent", str(-emu))
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(color)}))
    for el in (clr, pPr.makeelement(qn("a:buSzPct"), {"val": "70000"}),
               pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}),
               pPr.makeelement(qn("a:buChar"), {"char": char})):
        pPr.append(el)


# ---------------------------------------------------------------- shapes

def _shape(slide, kind, x, y, w, h, fill=WHITE, line=None, lw=0.75, radius=None, dash=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = 4
    if radius is not None and len(s.adjustments):
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, fill=WHITE, line=BORDER, lw=0.75, radius=0.09, dash=False):
    return _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, lw, radius, dash)


def oval(slide, x, y, w, h, fill, line=None):
    return _shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill, line)


def title_card(slide, x, y, w, h, title, size=12.5):
    card(slide, x, y, w, h)
    TX(slide, x, y + 0.07, w, 0.3, [[(title, {"bold": True, "color": BLUE})]], size=size,
       align=PP_ALIGN.CENTER)


def bar_card(slide, x, y, w, h, title, bar=0.34, size=12.5):
    """White bordered card with a blue header bar."""
    card(slide, x, y, w, h)
    _shape(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, y, w, bar, BLUE, BORDER, 0.75, 0.09)
    TX(slide, x, y, w, bar, [[(title, {"bold": True, "color": WHITE})]], size=size,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def line(slide, x1, y1, x2, y2, color=RULE, width=0.75, dash=False, arrow=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    if arrow in ("end", "both"):
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if arrow == "both":
        ln.insert(list(ln).index(ln.find(qn("a:tailEnd"))),
                  ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def chips(slide, x, y, w, items, size=8.5, h=0.225, gap=0.05, fill=WHITE, border=CHIP):
    """Flow pill chips inside width w; returns the y below the last row."""
    cx, cy = x, y
    for it in items:
        cw = 0.062 * len(it) * size / 8.5 + 0.16
        if cx + cw > x + w + 0.001 and cx > x:
            cx, cy = x, cy + h + gap
        s = card(slide, cx, cy, cw, h, fill=fill, line=border, lw=0.6, radius=0.06)
        tf = s.text_frame
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = it
        r.font.name, r.font.size, r.font.color.rgb = F, Pt(size), INK
        cx += cw + gap
    return cy + h


def icon_rows(slide, x, y, w, rows, row_h, size=10, icon=0.36, rule=True):
    """rows: (icon, color, text-with-**bold**)."""
    for i, (name, col, text) in enumerate(rows):
        ry = y + i * row_h
        ico(slide, name, col, x + 0.1, ry + (row_h - icon) / 2, icon)
        TX(slide, x + 0.1 + icon + 0.12, ry, w - icon - 0.32, row_h, text, size=size,
           anchor=MSO_ANCHOR.MIDDLE, space_after=0, line=0.95)
        if rule and i < len(rows) - 1:
            line(slide, x + 0.12, ry + row_h, x + w - 0.12, ry + row_h)


def gradient_run(run, c1, c2):
    rPr = run._r.get_or_add_rPr()
    for el in rPr.findall(qn("a:solidFill")):
        rPr.remove(el)
    g = rPr.makeelement(qn("a:gradFill"), {})
    lst = g.makeelement(qn("a:gsLst"), {})
    for pos, c in ((0, c1), (100000, c2)):
        gs = lst.makeelement(qn("a:gs"), {"pos": str(pos)})
        gs.append(gs.makeelement(qn("a:srgbClr"), {"val": c}))
        lst.append(gs)
    g.append(lst)
    g.append(g.makeelement(qn("a:lin"), {"ang": "0", "scaled": "0"}))
    rPr.insert(0, g)


def shot(slide, path, x, y, w, alt, h=None, border=RGBColor(0x9C, 0xA3, 0xAF)):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h) if h else None)
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    pic.line.color.rgb = border
    pic.line.width = Pt(0.75)
    return pic



